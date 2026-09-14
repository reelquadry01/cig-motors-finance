"""Generate blank Excel templates on demand.

Each template has:
  1. A data sheet with the correct headers and a single greyed-out sample row
     showing the expected format.
  2. A validation sheet with per-column notes so someone new to the schema
     can fill it in without asking.

Kept in Python (not files on disk) so the templates stay in lock-step with
FILE_TYPES in config.py. Change the spec, next download is up-to-date.
"""
from __future__ import annotations
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .. import config


BRAND = "C8102E"
NAVY = "1F3A5F"
MUTED = "7A736C"
HEADFILL = "F2F0ED"
RULE = "B9B2AA"


# Human-readable notes shown on the "Notes" sheet of each template.
COLUMN_NOTES: dict[str, str] = {
    "Transaction_ID": "Optional. If present, we use it as the primary diff key. Composite of GL_Code + Doc_Date + Reference is used otherwise.",
    "GL_Code": "Chart-of-accounts code. Must match a row in Statement_Mapping.xlsx.",
    "GL_Account": "Human-readable account name (used on the face of the statements).",
    "Doc_Date": "Transaction date in YYYY-MM-DD.",
    "Source": "Ledger source (Sales / Purchases / Cash / General …).",
    "Reference": "Invoice/receipt/journal number.",
    "Narration": "Free-text description.",
    "Debit": "Debit amount in Naira. Zero if credit-only.",
    "Credit": "Credit amount in Naira. Zero if debit-only.",
    "Account_Description": "Name of the GL account.",
    "Statement_Section": "Assets / Liabilities / Equity / Revenue / COGS / Operating Expenses / Depreciation / Other Income / Finance Costs / Tax.",
    "FS_Heading": "The line the account rolls up to on the face of the statements (e.g. \"Trade receivables\").",
    "Note_Heading": "The note the account rolls up into (e.g. \"Motor vehicles\").",
    "CF_Category": "Operating / Investing / Financing.",
    "Segment": "Motor Vehicles Sales / Spare Parts & After-Sales / Corporate.",
    "Borrowing_Type": "For borrowings only — Short-term / Long-term.",
    "Normal_Balance": "\"Debit\" (assets, expenses) or \"Credit\" (liabilities, equity, income).",
    "Period": "Reporting period key, YYYY-MM.",
    "Budget_Amount": "Budgeted figure for the account × period.",
    "Account_Name": "Human-readable account name.",
    "Opening_Balance_Dr": "Opening debit balance (Naira). Zero if credit-natured.",
    "Opening_Balance_Cr": "Opening credit balance (Naira). Zero if debit-natured.",
}


def _header_cell(cell) -> None:
    cell.font = Font(name="Aptos Narrow", bold=True, size=10, color="1A1A1A")
    cell.fill = PatternFill(fill_type="solid", start_color=HEADFILL, end_color=HEADFILL)
    cell.border = Border(bottom=Side(style="thin", color=RULE))
    cell.alignment = Alignment(vertical="center")


def _sample_cell(cell) -> None:
    cell.font = Font(name="Aptos Narrow", italic=True, size=10, color=MUTED)
    cell.alignment = Alignment(vertical="center")


def build_template(file_type: str) -> bytes:
    """Return a fully-styled .xlsx template for `file_type`."""
    spec = config.file_type_or_400(file_type)
    wb = Workbook()

    # Data sheet
    ws = wb.active
    ws.title = str(spec.get("sheet", "Data"))[:31]
    for i, h in enumerate(spec["headers"], start=1):
        c = ws.cell(row=1, column=i, value=h)
        _header_cell(c)
    for i, v in enumerate(spec["sample"], start=1):
        c = ws.cell(row=2, column=i, value=v)
        _sample_cell(c)
    for i, h in enumerate(spec["headers"], start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(28, len(h) + 6))
    ws.freeze_panes = "A2"

    # Notes sheet
    notes = wb.create_sheet("Notes")
    # Title
    t = notes.cell(row=1, column=1, value=f"{spec['label']} — column guide")
    t.font = Font(name="Aptos Narrow", bold=True, size=14, color=BRAND)
    notes.row_dimensions[1].height = 22

    subtitle = notes.cell(row=2, column=1, value="Fill the sheet named " + ws.title + ". Delete the italic sample row before uploading.")
    subtitle.font = Font(name="Aptos Narrow", italic=True, size=9, color=MUTED)

    # Header row
    for i, label in enumerate(["Column", "Meaning"], start=1):
        _header_cell(notes.cell(row=4, column=i, value=label))
    for r, h in enumerate(spec["headers"], start=5):
        notes.cell(row=r, column=1, value=h).font = Font(name="Aptos Narrow", bold=True, size=10, color=NAVY)
        cell = notes.cell(row=r, column=2, value=COLUMN_NOTES.get(h, "—"))
        cell.font = Font(name="Aptos Narrow", size=10, color="1A1A1A")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[r].height = max(18, 15 * (len(cell.value) // 90 + 1))
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 90
    notes.freeze_panes = "A5"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
