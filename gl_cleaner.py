"""
gl_cleaner.py
=============
Cleans a hierarchical General Ledger (GL) export into a flat, tidy,
analysis-ready transactions table, and reconciles it back to the source.

Built for CIG Motors monthly GL dumps, but written to be reusable for any
GL export that follows the same nested layout:

    [ACCOUNT HEADER]  code | account name | "Opening Balance:" | opening_dr | opening_cr
    [YEAR ROW]        2025 | (blank ...)
    [TRANSACTIONS]    period(01-12) | source | doc_date | narration | reference | debit | credit
    ...
    [TOTALS ROW]      "Totals: <name> <year>" | total_dr | total_cr
    [ENDING BALANCE]  "Ending Balance:"       | ending_dr | ending_cr

Outputs a workbook with three sheets:
  GL_Clean         one row per journal line, filled down and cleaned (values)
  Account_Summary  one row per account, fully formula-driven, reconciled
  Exceptions       anything needing review (unmapped codes, recon breaks)
plus a CSV of the clean transactions.

The Account_Summary is live: transaction debits, credits and counts are
pulled from GL_Clean with SUMIFS / COUNTIFS, and every net, movement,
closing and reconciliation check is an Excel formula. Opening balances and
the source Totals / Ending Balance figures are the only hardcoded inputs,
because they are the independent reference the sheet reconciles against.
The workbook is set to recalculate on open, so the formulas compute the
moment it is opened in Excel.

Account names (the agreed chart of accounts)
--------------------------------------------
The script needs a code-to-name chart to correct account descriptions. It can
read that chart from three places, in priority order:
  1. a file passed with --map-file
  2. a fixed file set once in DEFAULT_MAP_FILE at the top of this script
  3. a sheet inside the GL workbook itself (the original behaviour)
So you can keep one TB / chart file in a folder and run each month with only
the GL file. If no chart is found anywhere, the run still succeeds using the
GL's own account names, and every account is flagged Unmapped-Review.

Usage
-----
    python gl_cleaner.py INPUT.xlsx
    python gl_cleaner.py INPUT.xlsx -o OUTPUT.xlsx
    python gl_cleaner.py INPUT.xlsx --map-file "C:\\Finance\\Agreed_Chart.xlsx"
    python gl_cleaner.py INPUT.xlsx --gl-sheet gl --map-sheet "TB"

Requires: pandas, openpyxl.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule


# ----------------------------------------------------------------------------
# Configuration - safe defaults for the CIG GL layout. Override via CLI flags.
# ----------------------------------------------------------------------------
DEFAULT_GL_SHEET = "gl"
DEFAULT_MAP_SHEET = "correct full TB codes & Desc"

# Optional: a fixed chart-of-accounts (TB) file to read account names from,
# so each monthly GL export only needs the ledger sheet. Set the path once
# here and you can then run the script with just the GL file. Leave it as an
# empty string to disable. The --map-file flag overrides this at run time.
# Windows example: DEFAULT_MAP_FILE = r"C:\Finance\Reference\Agreed_Chart.xlsx"
DEFAULT_MAP_FILE = r""

# Sheet names to try, in order, when the preferred map sheet is not found.
MAP_SHEET_FALLBACKS = [
    "correct full TB codes & Desc", "TB", "Chart", "Chart of Accounts", "COA",
]

MONTH_NAMES = {
    1: "January", 2: "February", 3: "March", 4: "April",
    5: "May", 6: "June", 7: "July", 8: "August",
    9: "September", 10: "October", 11: "November", 12: "December",
}

RECON_TOLERANCE = 0.5  # currency rounding tolerance for reconciliation

# GL_Clean column order (also fixes the SUMIFS references in the summary).
GLCLEAN_COLS = [
    "GL_Code", "GL_Account", "Desc_Status", "Doc_Date", "Year", "Month_No",
    "Month_Name", "Period", "Source", "Reference", "Narration",
    "Debit", "Credit", "Net", "Source_Account_Name",
]
# Column letters within GL_Clean used by the summary formulas.
CLEAN_CODE_COL = "A"    # GL_Code
CLEAN_DEBIT_COL = "L"   # Debit
CLEAN_CREDIT_COL = "M"  # Credit


# ----------------------------------------------------------------------------
# Small helpers
# ----------------------------------------------------------------------------
def clean_text(value) -> str:
    """Trim, collapse internal whitespace, and strip control characters.

    Business narrations and references keep their original spelling and case;
    only formatting noise is removed.
    """
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[\x00-\x1f\x7f]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def to_number(value) -> float:
    """Coerce a cell to a float. Blanks and non-numeric text become 0.0."""
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if text in ("", "-"):
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def classify_row(row) -> str:
    """Return the structural type of a raw GL row based on its first cell."""
    first = row[0]
    text = "" if first is None else str(first).strip()
    if text == "":
        return "BLANK"
    lower = text.lower()
    if lower.startswith("account number"):
        return "FILE_HEADER"
    if lower.startswith("totals:"):
        return "TOTALS"
    if lower.startswith("report totals"):
        return "REPORT_TOTALS"
    if lower.startswith("ending balance"):
        return "ENDING"
    if "accounts printed" in lower or lower.startswith("net income"):
        return "FOOTER"
    if text.isdigit():
        if len(text) == 4 and text.startswith("20"):
            return "YEAR"
        if len(text) == 2 and 1 <= int(text) <= 12:
            return "PERIOD"
        if len(text) == 5:
            return "ACCOUNT_HEADER"
    return "OTHER"


# ----------------------------------------------------------------------------
# Reference data
# ----------------------------------------------------------------------------
def _is_code_like(value) -> bool:
    """True if the value looks like a 4 to 6 digit GL account code."""
    if value is None:
        return False
    text = str(value).strip()
    return text.isdigit() and 4 <= len(text) <= 6


def load_agreed_map_from_ws(ws) -> dict:
    """Build {gl_code: description} from a worksheet.

    The code and description columns are detected per row rather than assumed,
    so a plain two-column list, a full trial balance with extra columns, and a
    file with header or footer rows all work. Any row without a code-like cell
    (headers, totals, "accounts printed" footers) is skipped.
    """
    mapping = {}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        code = None
        code_idx = None
        for idx, cell in enumerate(row):
            if _is_code_like(cell):
                code = str(cell).strip()
                code_idx = idx
                break
        if code is None:
            continue
        desc = ""
        for cell in row[code_idx + 1:]:
            text = clean_text(cell)
            if text and not _is_code_like(cell):
                desc = text
                break
        mapping.setdefault(code, desc)  # first occurrence wins
    return mapping


def _looks_like_ledger(ws) -> bool:
    """True if a sheet is a hierarchical GL ledger, not a chart of accounts.

    Guards against ever scraping account codes out of the ledger itself.
    """
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 40:
            break
        for cell in row[:4]:
            if cell is None:
                continue
            text = str(cell).strip().lower()
            if (text.startswith("opening balance")
                    or text.startswith("ending balance")
                    or text.startswith("account number/year")):
                return True
    return False


def _pick_map_sheet(workbook, preferred: str, exclude=(), allow_first=False):
    """Choose a chart sheet, tolerant of naming and never a ledger sheet.

    Returns the sheet name, or None if no plausible chart sheet exists.
    """
    excl = {e.strip().lower() for e in exclude}
    candidates = [
        n for n in workbook.sheetnames
        if n.strip().lower() not in excl and not _looks_like_ledger(workbook[n])
    ]
    if not candidates:
        return None
    lowered = {n.strip().lower(): n for n in candidates}
    if preferred.strip().lower() in lowered:
        return lowered[preferred.strip().lower()]
    for candidate in MAP_SHEET_FALLBACKS:
        if candidate.strip().lower() in lowered:
            return lowered[candidate.strip().lower()]
    return candidates[0] if allow_first else None


def resolve_map(gl_path: str, gl_sheet: str, map_file: str, map_sheet: str):
    """Resolve the agreed chart of accounts and return (mapping, info_string).

    Source priority: --map-file, then DEFAULT_MAP_FILE, then a chart sheet
    inside the GL workbook. A named-but-missing external file falls back to the
    GL workbook. The ledger sheet is never read as a chart. If no chart is found
    anywhere, the map is empty and the run still succeeds using the GL's own
    names, with every account flagged Unmapped-Review.
    """
    chosen = (map_file or DEFAULT_MAP_FILE or "").strip() or None
    external = chosen is not None
    if chosen is not None and not Path(chosen).exists():
        print(f"WARNING: chart file not found, looking inside the GL instead: {chosen}",
              file=sys.stderr)
        chosen = None
        external = False
    source_path = Path(chosen) if chosen else Path(gl_path)

    wb = openpyxl.load_workbook(source_path, data_only=True)
    if external:
        # The user pointed us at a chart file on purpose: trust its first
        # non-ledger sheet if nothing better matches.
        sheet = _pick_map_sheet(wb, map_sheet, allow_first=True)
    else:
        # Internal fallback: only a real chart sheet, never the ledger.
        sheet = _pick_map_sheet(wb, map_sheet, exclude=(gl_sheet,), allow_first=False)

    mapping = load_agreed_map_from_ws(wb[sheet]) if sheet else {}
    origin = "external chart" if external else "GL workbook"
    where = f"sheet '{sheet}' in {source_path.name}" if sheet else f"none found in {source_path.name}"
    info = f"{len(mapping)} codes from {where} ({origin})"
    if not mapping:
        print("WARNING: no chart of accounts found. Account names will be taken "
              "from the GL as-is and every account flagged Unmapped-Review. "
              "Provide one with --map-file or set DEFAULT_MAP_FILE.",
              file=sys.stderr)
    return mapping, info


# ----------------------------------------------------------------------------
# Core parser
# ----------------------------------------------------------------------------
def parse_gl(path: str, gl_sheet: str, agreed: dict):
    """Parse a GL workbook into transactions, an account summary, and exceptions.

    `agreed` is the {code: description} chart resolved by resolve_map().
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    if gl_sheet not in wb.sheetnames:
        raise KeyError(f"GL sheet '{gl_sheet}' not found. Available: {wb.sheetnames}")

    ws = wb[gl_sheet]

    transactions = []
    accounts = {}
    order = []
    current = None

    for row in ws.iter_rows(values_only=True):
        kind = classify_row(row)

        if kind == "ACCOUNT_HEADER":
            code = str(row[0]).strip()
            raw_name = clean_text(row[1])
            agreed_name = agreed.get(code)
            matched = agreed_name is not None
            current = {
                "GL_Code": code,
                "Source_Account_Name": raw_name,
                "GL_Account": agreed_name if matched else raw_name,
                "Desc_Status": "Matched" if matched else "Unmapped-Review",
                "Opening_Debit": to_number(row[3]),
                "Opening_Credit": to_number(row[4]),
                "Txn_Debit": 0.0,
                "Txn_Credit": 0.0,
                "Txn_Count": 0,
                "Src_Total_Debit": None,
                "Src_Total_Credit": None,
                "Src_Ending_Debit": None,
                "Src_Ending_Credit": None,
            }
            if code not in accounts:
                accounts[code] = current
                order.append(code)
            else:
                current = accounts[code]

        elif kind == "PERIOD" and current is not None:
            doc_date = row[2]
            debit = to_number(row[5])
            credit = to_number(row[6])
            year = getattr(doc_date, "year", None)
            month = getattr(doc_date, "month", None)
            transactions.append({
                "GL_Code": current["GL_Code"],
                "GL_Account": current["GL_Account"],
                "Desc_Status": current["Desc_Status"],
                "Doc_Date": doc_date.date() if hasattr(doc_date, "date") else doc_date,
                "Year": year,
                "Month_No": month,
                "Month_Name": MONTH_NAMES.get(month, ""),
                "Period": f"{year}-{month:02d}" if year and month else "",
                "Source": clean_text(row[1]),
                "Reference": clean_text(row[4]),
                "Narration": clean_text(row[3]),
                "Debit": debit,
                "Credit": credit,
                "Net": round(debit - credit, 2),
                "Source_Account_Name": current["Source_Account_Name"],
            })
            current["Txn_Debit"] += debit
            current["Txn_Credit"] += credit
            current["Txn_Count"] += 1

        elif kind == "TOTALS" and current is not None:
            current["Src_Total_Debit"] = to_number(row[1])
            current["Src_Total_Credit"] = to_number(row[2])

        elif kind == "ENDING" and current is not None:
            current["Src_Ending_Debit"] = to_number(row[1])
            current["Src_Ending_Credit"] = to_number(row[2])

    # Account summary (values used for the Exceptions log and console summary;
    # the workbook itself re-derives these with live formulas).
    summary = []
    exceptions = []
    for code in order:
        a = accounts[code]
        opening_net = round(a["Opening_Debit"] - a["Opening_Credit"], 2)
        movement = round(a["Txn_Debit"] - a["Txn_Credit"], 2)
        computed_closing = round(opening_net + movement, 2)
        src_total_dr = a["Src_Total_Debit"]
        src_total_cr = a["Src_Total_Credit"]
        src_ending_net = None
        if a["Src_Ending_Debit"] is not None:
            src_ending_net = round(a["Src_Ending_Debit"] - a["Src_Ending_Credit"], 2)

        totals_ok = True
        if src_total_dr is not None:
            totals_ok = (
                abs(a["Txn_Debit"] - src_total_dr) <= RECON_TOLERANCE
                and abs(a["Txn_Credit"] - src_total_cr) <= RECON_TOLERANCE
            )
        ending_ok = True
        if src_ending_net is not None:
            ending_ok = abs(computed_closing - src_ending_net) <= RECON_TOLERANCE

        summary.append({
            "GL_Code": code,
            "GL_Account": a["GL_Account"],
            "Desc_Status": a["Desc_Status"],
            "Opening_Debit": a["Opening_Debit"],
            "Opening_Credit": a["Opening_Credit"],
            "Src_Total_Debit": src_total_dr,
            "Src_Total_Credit": src_total_cr,
            "Src_Ending_Net": src_ending_net,
            # value-form derivations retained for the console summary only
            "_movement": movement,
            "_closing": computed_closing,
            "_txn_count": a["Txn_Count"],
        })

        if a["Desc_Status"] != "Matched":
            exceptions.append({
                "GL_Code": code,
                "Issue": "GL code not in agreed chart of accounts",
                "Detail": f"Kept source name: {a['Source_Account_Name']}",
            })
        if not totals_ok:
            exceptions.append({
                "GL_Code": code,
                "Issue": "Transaction totals do not match source Totals row",
                "Detail": (
                    f"clean dr/cr {a['Txn_Debit']:.2f}/{a['Txn_Credit']:.2f} "
                    f"vs source {src_total_dr:.2f}/{src_total_cr:.2f}"
                ),
            })
        if not ending_ok:
            exceptions.append({
                "GL_Code": code,
                "Issue": "Computed closing does not match source Ending Balance",
                "Detail": f"computed {computed_closing:.2f} vs source {src_ending_net:.2f}",
            })

    return transactions, summary, exceptions


# ----------------------------------------------------------------------------
# Workbook construction (formula-driven summary + formatting)
# ----------------------------------------------------------------------------
ARIAL = "Arial"
MONEY_FMT = "#,##0;(#,##0);-"
COUNT_FMT = "#,##0"
DATE_FMT = "yyyy-mm-dd"

SUMMARY_HEADERS = [
    "GL_Code", "GL_Account", "Desc_Status",
    "Opening_Debit", "Opening_Credit", "Opening_Net",
    "Txn_Debit", "Txn_Credit", "Movement", "Computed_Closing_Net",
    "Src_Total_Debit", "Src_Total_Credit", "Src_Ending_Net",
    "Txn_Count", "Recon_Totals_OK", "Recon_Ending_OK",
]


def _build_account_summary(wb, summary):
    """Create Account_Summary with hardcoded inputs and live formulas."""
    ws = wb.create_sheet("Account_Summary", 1)
    ws.append(SUMMARY_HEADERS)

    for i, a in enumerate(summary):
        r = i + 2
        ws.cell(r, 1, a["GL_Code"])
        ws.cell(r, 2, a["GL_Account"])
        ws.cell(r, 3, a["Desc_Status"])
        ws.cell(r, 4, a["Opening_Debit"])            # input
        ws.cell(r, 5, a["Opening_Credit"])           # input
        ws.cell(r, 6, f"=D{r}-E{r}")                 # Opening_Net
        ws.cell(r, 7, f"=SUMIFS(GL_Clean!${CLEAN_DEBIT_COL}:${CLEAN_DEBIT_COL},"
                      f"GL_Clean!${CLEAN_CODE_COL}:${CLEAN_CODE_COL},$A{r})")   # Txn_Debit
        ws.cell(r, 8, f"=SUMIFS(GL_Clean!${CLEAN_CREDIT_COL}:${CLEAN_CREDIT_COL},"
                      f"GL_Clean!${CLEAN_CODE_COL}:${CLEAN_CODE_COL},$A{r})")   # Txn_Credit
        ws.cell(r, 9, f"=G{r}-H{r}")                 # Movement
        ws.cell(r, 10, f"=F{r}+I{r}")                # Computed_Closing_Net
        ws.cell(r, 11, a["Src_Total_Debit"])         # input
        ws.cell(r, 12, a["Src_Total_Credit"])        # input
        ws.cell(r, 13, a["Src_Ending_Net"])          # input
        ws.cell(r, 14, f"=COUNTIFS(GL_Clean!${CLEAN_CODE_COL}:${CLEAN_CODE_COL},$A{r})")  # Txn_Count
        ws.cell(r, 15, f'=IF(AND(ABS(G{r}-K{r})<=0.5,ABS(H{r}-L{r})<=0.5),"OK","CHECK")')
        ws.cell(r, 16, f'=IF(ABS(J{r}-M{r})<=0.5,"OK","CHECK")')

    total_r = len(summary) + 2
    ws.cell(total_r, 1, "TOTAL")
    for col in (4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14):
        letter = get_column_letter(col)
        ws.cell(total_r, col, f"=SUM({letter}2:{letter}{total_r - 1})")
    ws.cell(total_r, 15, f'=COUNTIF(O2:O{total_r - 1},"CHECK")')
    ws.cell(total_r, 16, f'=COUNTIF(P2:P{total_r - 1},"CHECK")')
    return ws, total_r


def _style_headers(ws, ncols):
    hdr_font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    for c in range(1, ncols + 1):
        cell = ws.cell(1, c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _apply_base_font(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.font is None or cell.font.name != ARIAL:
                cell.font = Font(name=ARIAL, size=10)


def _format_gl_clean(ws, n_tx):
    _apply_base_font(ws)
    _style_headers(ws, len(GLCLEAN_COLS))
    last = n_tx + 1
    for r in range(2, last + 1):
        ws.cell(r, 4).number_format = DATE_FMT           # Doc_Date
        for c in (12, 13, 14):                           # Debit, Credit, Net
            ws.cell(r, c).number_format = MONEY_FMT
    widths = {"A": 10, "B": 34, "C": 15, "D": 12, "E": 6, "F": 8, "G": 11,
              "H": 9, "I": 8, "J": 20, "K": 44, "L": 16, "M": 16, "N": 16, "O": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def _format_summary(ws, n_sum, total_r):
    _apply_base_font(ws)
    _style_headers(ws, len(SUMMARY_HEADERS))

    blue = Font(name=ARIAL, size=10, color="0000FF")   # hardcoded input
    green = Font(name=ARIAL, size=10, color="008000")  # link to GL_Clean
    black = Font(name=ARIAL, size=10, color="000000")  # in-sheet formula
    red = Font(name=ARIAL, size=10, color="FF0000")

    input_cols = (4, 5, 11, 12, 13)      # opening + source totals/ending
    link_cols = (7, 8, 14)               # SUMIFS / COUNTIFS from GL_Clean
    formula_cols = (6, 9, 10, 15, 16)    # in-sheet derivations
    money_cols = (4, 5, 6, 7, 8, 9, 10, 11, 12, 13)

    for r in range(2, n_sum + 2):
        for c in input_cols:
            ws.cell(r, c).font = blue
        for c in link_cols:
            ws.cell(r, c).font = green
        for c in formula_cols:
            ws.cell(r, c).font = black
        for c in money_cols:
            ws.cell(r, c).number_format = MONEY_FMT
        ws.cell(r, 14).number_format = COUNT_FMT
        if ws.cell(r, 3).value and ws.cell(r, 3).value != "Matched":
            ws.cell(r, 3).font = red

    # TOTAL row
    bold = Font(name=ARIAL, size=10, bold=True)
    topb = Border(top=Side(style="thin", color="000000"))
    for c in range(1, len(SUMMARY_HEADERS) + 1):
        cell = ws.cell(total_r, c)
        cell.font = bold
        cell.border = topb
        if c in money_cols:
            cell.number_format = MONEY_FMT
        if c == 14:
            cell.number_format = COUNT_FMT

    # Highlight any reconciliation break in red.
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    ws.conditional_formatting.add(
        f"O2:P{n_sum + 1}",
        CellIsRule(operator="equal", formula=['"CHECK"'],
                   fill=red_fill, font=Font(name=ARIAL, color="9C0006")),
    )

    widths = {"A": 10, "B": 34, "C": 16, "D": 16, "E": 16, "F": 16, "G": 16,
              "H": 16, "I": 16, "J": 18, "K": 16, "L": 16, "M": 16,
              "N": 11, "O": 15, "P": 15}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "D2"

    # Legend for the colour convention, two rows below the TOTAL row.
    lr = total_r + 2
    ws.cell(lr, 1, "Legend").font = bold
    ws.cell(lr + 1, 1, "Blue: source input from the raw GL").font = blue
    ws.cell(lr + 2, 1, "Green: live link to GL_Clean (SUMIFS / COUNTIFS)").font = green
    ws.cell(lr + 3, 1, "Black: calculated in this sheet").font = black
    ws.cell(lr + 4, 1,
            'Recon columns show OK, or CHECK if the account does not tie to source.').font = \
        Font(name=ARIAL, size=10, italic=True)


def _format_exceptions(ws):
    _apply_base_font(ws)
    _style_headers(ws, 3)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 60
    ws.freeze_panes = "A2"


def write_outputs(transactions, summary, exceptions, out_xlsx: Path):
    tx = pd.DataFrame(transactions, columns=GLCLEAN_COLS)
    if not tx.empty:
        tx = tx.sort_values(["GL_Code", "Doc_Date"]).reset_index(drop=True)
    ex = pd.DataFrame(exceptions) if exceptions else pd.DataFrame(
        columns=["GL_Code", "Issue", "Detail"]
    )

    # Bulk sheets via pandas.
    with pd.ExcelWriter(out_xlsx, engine="openpyxl") as writer:
        tx.to_excel(writer, sheet_name="GL_Clean", index=False)
        ex.to_excel(writer, sheet_name="Exceptions", index=False)

    # Formula-driven summary + formatting via openpyxl.
    wb = load_workbook(out_xlsx)
    _, total_r = _build_account_summary(wb, summary)
    _format_gl_clean(wb["GL_Clean"], n_tx=len(tx))
    _format_summary(wb["Account_Summary"], n_sum=len(summary), total_r=total_r)
    _format_exceptions(wb["Exceptions"])
    wb.calculation.fullCalcOnLoad = True   # recompute on open in Excel
    wb.save(out_xlsx)

    csv_path = out_xlsx.with_name(out_xlsx.stem + "_transactions.csv")
    tx.to_csv(csv_path, index=False)
    return tx, pd.DataFrame(summary), ex, csv_path


# ----------------------------------------------------------------------------
# Console summary
# ----------------------------------------------------------------------------
def print_summary(tx, summary, ex, csv_path, out_xlsx, map_info=""):
    line = "-" * 60
    print(line)
    print("GL CLEAN - RUN SUMMARY")
    print(line)
    if map_info:
        print(f"Chart of accounts          : {map_info}")
    print(f"Transactions (clean lines) : {len(tx):,}")
    print(f"Accounts                   : {len(summary):,}")
    if not tx.empty:
        print(f"Total debits               : {tx['Debit'].sum():,.2f}")
        print(f"Total credits              : {tx['Credit'].sum():,.2f}")
        print(f"Net movement               : {tx['Net'].sum():,.2f}")
        print(f"Date range                 : {tx['Doc_Date'].min()} to {tx['Doc_Date'].max()}")
    unmatched = sum(1 for a in summary if a["Desc_Status"] != "Matched")
    print(f"Unmapped account codes     : {unmatched}")
    print(f"Exceptions logged          : {len(ex)}")
    print(line)
    print(f"Workbook : {out_xlsx}")
    print(f"CSV      : {csv_path}")
    print(line)


# ----------------------------------------------------------------------------
# CLI
# ----------------------------------------------------------------------------
def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Clean a hierarchical GL export into a flat, reconciled table."
    )
    parser.add_argument("input", help="Path to the GL .xlsx export")
    parser.add_argument("-o", "--output", help="Output .xlsx (default: <input>_clean.xlsx)")
    parser.add_argument("--gl-sheet", default=DEFAULT_GL_SHEET,
                        help=f"Ledger sheet name in the GL file (default: {DEFAULT_GL_SHEET})")
    parser.add_argument("--map-file", default=None,
                        help="Path to a separate chart-of-accounts / TB file to read "
                             "account names from. Overrides DEFAULT_MAP_FILE.")
    parser.add_argument("--map-sheet", default=DEFAULT_MAP_SHEET,
                        help="Sheet name that holds the code and description "
                             f"(default: '{DEFAULT_MAP_SHEET}').")
    args = parser.parse_args(argv)

    in_path = Path(args.input)
    if not in_path.exists():
        print(f"Input not found: {in_path}", file=sys.stderr)
        return 1
    out_xlsx = Path(args.output) if args.output else in_path.with_name(
        in_path.stem + "_clean.xlsx"
    )

    agreed, map_info = resolve_map(str(in_path), args.gl_sheet, args.map_file, args.map_sheet)
    transactions, summary, exceptions = parse_gl(str(in_path), args.gl_sheet, agreed)
    tx, sm, ex, csv_path = write_outputs(transactions, summary, exceptions, out_xlsx)
    print_summary(tx, summary, ex, csv_path, out_xlsx, map_info)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
